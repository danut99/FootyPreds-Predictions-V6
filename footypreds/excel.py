"""Daily predictions workbook (.xlsx): one sheet per market family, like a desktop tool."""

import io
from datetime import timezone

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from footypreds.sports import headline_tip

HEADER_FILL = PatternFill("solid", fgColor="16211D")
HEADER_FONT = Font(bold=True, color="E8FF9C")
GRADE_FILLS = {
    "A": PatternFill("solid", fgColor="C6EFCE"),
    "B": PatternFill("solid", fgColor="E2F0D9"),
    "C": PatternFill("solid", fgColor="FFF2CC"),
    "D": PatternFill("solid", fgColor="F8CBAD"),
}
PERCENT = "0%"


def probabilities(analysis):
    return {m["key"]: m for m in analysis["markets"]}


def clean(value):
    """Feed text without XML-illegal control characters (they make openpyxl raise)."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def append_row(ws, row):
    """Append values as data: a string is always text, never a formula (e.g. '=HYPERLINK')."""
    values = [clean(value) for value in row]
    ws.append(values)
    for column, value in enumerate(values, 1):
        if isinstance(value, str):
            ws.cell(row=ws.max_row, column=column).data_type = "s"


def sheet(workbook, title, headers, rows, percent_columns=(), widths=None, scale=()):
    ws = workbook.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        append_row(ws, row)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for index, header in enumerate(headers, 1):
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = (widths or {}).get(header, max(9, len(header) + 2))
        if header in percent_columns:
            for (cell,) in ws.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = PERCENT
        if header in scale and rows:
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{len(rows) + 1}",
                ColorScaleRule(
                    start_type="num",
                    start_value=0,
                    start_color="FFFFFF",
                    end_type="num",
                    end_value=1,
                    end_color="63BE7B",
                ),
            )
    return ws


def local_time(match):
    return match.kickoff.astimezone(timezone.utc).strftime("%H:%M")


def build_workbook(day, items, source="FlashScore"):
    """items: list of (match, analysis). Returns .xlsx bytes."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    main_headers = [
        "Ora (UTC)",
        "Competiție",
        "Gazde",
        "Oaspeți",
        "1",
        "X",
        "2",
        "1X",
        "X2",
        "Peste 1.5",
        "Peste 2.5",
        "Sub 2.5",
        "Peste 3.5",
        "GG",
        "xG gazde",
        "xG oaspeți",
        "Scor probabil",
        "Pont principal",
        "Prob. pont",
        "Cotă 1",
        "Cotă X",
        "Cotă 2",
        "Calitate",
        "Încredere",
        "Formă gazde",
        "Formă oaspeți",
    ]
    main_rows, scores, form_rows, value_rows, ht_rows = [], [], [], [], []
    for match, analysis in items:
        p = probabilities(analysis)
        tip = headline_tip(analysis)
        league = match.league.split(":", 1)[-1].strip()
        main_rows.append(
            [
                local_time(match),
                league,
                match.home,
                match.away,
                *(
                    p[k]["probability"]
                    for k in ("1", "X", "2", "1X", "X2", "over15", "over25", "under25", "over35")
                ),
                p["btts"]["probability"],
                round(analysis["expected_goals"]["home"], 2),
                round(analysis["expected_goals"]["away"], 2),
                analysis["scores"][0]["score"],
                tip["label"],
                tip["probability"],
                match.odds.get("1"),
                match.odds.get("X"),
                match.odds.get("2"),
                analysis["grade"],
                analysis["confidence"],
                analysis["form"]["home"]["sequence"] or "—",
                analysis["form"]["away"]["sequence"] or "—",
            ]
        )
        scores.append(
            [
                local_time(match),
                match.home,
                match.away,
                *[x for s in analysis["scores"][:5] for x in (s["score"], s["probability"])],
            ]
        )
        for side, name in (("home", match.home), ("away", match.away)):
            form = analysis["form"][side]
            last = form["last10"] or {}
            venue = form["home10" if side == "home" else "away10"] or {}
            form_rows.append(
                [
                    f"{match.home} – {match.away}",
                    name,
                    "Gazde" if side == "home" else "Oaspeți",
                    form["sequence"] or "—",
                    last.get("played", 0),
                    last.get("points_per_game"),
                    last.get("scored_avg"),
                    last.get("conceded_avg"),
                    last.get("over25"),
                    last.get("btts"),
                    last.get("clean_sheets"),
                    last.get("failed_to_score"),
                    venue.get("points_per_game"),
                    form["streaks"]["unbeaten"],
                    form["days_since_last"],
                ]
            )
        for market in analysis["markets"]:
            if market["ev"] is not None and market["ev"] > 0:
                value_rows.append(
                    [
                        local_time(match),
                        match.home,
                        match.away,
                        market["label"],
                        market["probability"],
                        round(market["fair_odds"], 2),
                        market["odds"],
                        market["ev"],
                        analysis["grade"],
                    ]
                )
        htft = analysis["htft"][:3]
        ht_rows.append(
            [
                local_time(match),
                match.home,
                match.away,
                p["ht_1"]["probability"],
                p["ht_X"]["probability"],
                p["ht_2"]["probability"],
                p["ht_over05"]["probability"],
                p["ht_over15"]["probability"],
                *[x for item in htft for x in (item["key"], item["probability"])],
            ]
        )
    probability_columns = main_headers[4:14] + ["Prob. pont"]
    ws = sheet(
        workbook,
        "Predicții",
        main_headers,
        main_rows,
        percent_columns=probability_columns,
        widths={"Competiție": 26, "Gazde": 22, "Oaspeți": 22, "Pont principal": 22},
        scale=probability_columns,
    )
    grade_column = main_headers.index("Calitate") + 1
    for (cell,) in ws.iter_rows(min_row=2, min_col=grade_column, max_col=grade_column):
        cell.fill = GRADE_FILLS.get(cell.value, PatternFill())
        cell.alignment = Alignment(horizontal="center")
    score_headers = ["Ora (UTC)", "Gazde", "Oaspeți"]
    for n in range(1, 6):
        score_headers += [f"Scor {n}", f"Prob. {n}"]
    sheet(
        workbook,
        "Scor corect",
        score_headers,
        scores,
        percent_columns=[f"Prob. {n}" for n in range(1, 6)],
        widths={"Gazde": 22, "Oaspeți": 22},
    )
    form_headers = [
        "Meci",
        "Echipă",
        "Rol",
        "Ultimele 5",
        "Meciuri",
        "Puncte/meci",
        "Goluri marcate/meci",
        "Goluri primite/meci",
        "Peste 2.5",
        "GG",
        "Fără gol primit",
        "Fără gol marcat",
        "Puncte/meci acasă sau deplasare",
        "Serie fără înfrângere",
        "Zile de la ultimul meci",
    ]
    ws = sheet(
        workbook,
        "Formă",
        form_headers,
        form_rows,
        percent_columns=["Peste 2.5", "GG", "Fără gol primit", "Fără gol marcat"],
        widths={"Meci": 40, "Echipă": 22},
    )
    for column in ("Puncte/meci", "Goluri marcate/meci", "Goluri primite/meci"):
        index = form_headers.index(column) + 1
        for (cell,) in ws.iter_rows(min_row=2, min_col=index, max_col=index):
            cell.number_format = "0.00"
    sheet(
        workbook,
        "Valoare",
        [
            "Ora (UTC)",
            "Gazde",
            "Oaspeți",
            "Piață",
            "Probabilitate",
            "Cotă corectă",
            "Cotă",
            "EV",
            "Calitate",
        ],
        sorted(value_rows, key=lambda r: r[7], reverse=True),
        percent_columns=["Probabilitate", "EV"],
        widths={"Gazde": 22, "Oaspeți": 22, "Piață": 20},
    )
    ht_headers = [
        "Ora (UTC)",
        "Gazde",
        "Oaspeți",
        "Pauză 1",
        "Pauză X",
        "Pauză 2",
        "Pauză peste 0.5",
        "Pauză peste 1.5",
        "HT/FT 1",
        "Prob. 1",
        "HT/FT 2",
        "Prob. 2",
        "HT/FT 3",
        "Prob. 3",
    ]
    sheet(
        workbook,
        "Pauză-Final",
        ht_headers,
        ht_rows,
        percent_columns=[h for h in ht_headers if h.startswith(("Pauză", "Prob."))],
        widths={"Gazde": 22, "Oaspeți": 22},
    )
    legend = workbook.create_sheet("Legendă")
    for line in (
        [f"FootyPreds — predicții pentru {day.isoformat()}", ""],
        ["Sursă meciuri și cote", source],
        ["Model", items[0][1]["version"] if items else "—"],
        ["", ""],
        ["Calitate A/B/C/D", "Cât de multe și de recente sunt datele pentru ambele echipe."],
        ["Încredere", "0–100: istoric recent pe echipă (toate competițiile) și cote disponibile."],
        ["xG", "Goluri așteptate de model pentru fiecare echipă."],
        ["Cotă corectă", "1 / probabilitate. Sub cota casei = valoare pozitivă (EV > 0)."],
        ["EV", "Probabilitate × cotă − 1. Estimare, nu profit garantat."],
        ["Formă", "W = victorie, D = egal, L = înfrângere; cel mai recent meci primul."],
        ["", ""],
        ["Atenție", "Probabilitățile sunt estimări statistice, nu garanții. 18+."],
    ):
        append_row(legend, line)
    legend.column_dimensions["A"].width = 26
    legend.column_dimensions["B"].width = 80
    legend["A1"].font = Font(bold=True, size=14)
    workbook.move_sheet("Legendă", offset=-(len(workbook.sheetnames) - 1))
    workbook.active = 1
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_sport_workbook(day, items, sport, source="FlashScore"):
    """Workbook for basketball/tennis (common analysis shape). items: (match, analysis)."""
    from footypreds.sports import SPORTS, main_markets

    workbook = Workbook()
    workbook.remove(workbook.active)
    main_headers = [
        "Ora (UTC)",
        "Competiție",
        "Gazde",
        "Oaspeți",
        "Piețe principale",
        "Pont principal",
        "Prob. pont",
        "Cotă 1",
        "Cotă 2",
        "Calitate",
        "Încredere",
        "Formă gazde",
        "Formă oaspeți",
    ]
    main_rows, market_rows = [], []
    for match, analysis in items:
        tip = headline_tip(analysis)
        headline = "; ".join(f"{m['label']} {m['probability']:.0%}" for m in main_markets(analysis))
        main_rows.append(
            [
                local_time(match),
                match.league.split(":", 1)[-1].strip(),
                match.home,
                match.away,
                headline,
                tip["label"],
                tip["probability"],
                match.odds.get("1"),
                match.odds.get("2"),
                analysis["grade"],
                analysis["confidence"],
                analysis["form"]["home"]["sequence"] or "—",
                analysis["form"]["away"]["sequence"] or "—",
            ]
        )
        for market in analysis["markets"]:
            market_rows.append(
                [
                    local_time(match),
                    match.home,
                    match.away,
                    market["group"],
                    market["label"],
                    market["probability"],
                    round(market["fair_odds"], 2) if market["fair_odds"] else None,
                    market["odds"],
                    market["ev"],
                ]
            )
    ws = sheet(
        workbook,
        "Predicții",
        main_headers,
        main_rows,
        percent_columns=["Prob. pont"],
        widths={"Competiție": 26, "Gazde": 22, "Oaspeți": 22, "Piețe principale": 60},
        scale=["Prob. pont"],
    )
    grade_column = main_headers.index("Calitate") + 1
    for (cell,) in ws.iter_rows(min_row=2, min_col=grade_column, max_col=grade_column):
        cell.fill = GRADE_FILLS.get(cell.value, PatternFill())
    market_headers = [
        "Ora (UTC)",
        "Gazde",
        "Oaspeți",
        "Grup",
        "Piață",
        "Probabilitate",
        "Cotă corectă",
        "Cotă",
        "EV",
    ]
    sheet(
        workbook,
        "Piețe",
        market_headers,
        market_rows,
        percent_columns=["Probabilitate", "EV"],
        widths={"Gazde": 22, "Oaspeți": 22, "Piață": 28, "Grup": 20},
    )
    legend = workbook.create_sheet("Legendă")
    label = SPORTS.get(sport, {}).get("label", sport)
    for line in (
        [f"FootyPreds — {label}, predicții pentru {day.isoformat()}", ""],
        ["Sursă meciuri și cote", source],
        ["Model", items[0][1]["version"] if items else "—"],
        ["Cotă corectă", "1 / probabilitate. Sub cota casei = valoare pozitivă (EV > 0)."],
        ["Atenție", "Probabilitățile sunt estimări statistice, nu garanții. 18+."],
    ):
        append_row(legend, line)
    legend.column_dimensions["A"].width = 26
    legend.column_dimensions["B"].width = 80
    workbook.move_sheet("Legendă", offset=-(len(workbook.sheetnames) - 1))
    workbook.active = 1
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
