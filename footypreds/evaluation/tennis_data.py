"""tennis-data.co.uk ATP and WTA results with bookmaker odds: download, cache, parse.

Source: https://www.tennis-data.co.uk/alldata.php lists one workbook per tour and year. The
links sit under an opaque directory that the site may rename, so `download` reads the index
page first and falls back to the last known prefix. Raw workbooks are kept unchanged in
`data/benchmark/tennis/raw/{tour}_{year}.xlsx` (git-ignored), with a manifest of URL, size and
SHA256. Only `.xlsx` files (2013 onwards) are supported: older years are `.xls`.

Parsing (`parse_workbook`) turns every row into a tennis `Match`:

- `home`/`away`: winner and loser, swapped by a deterministic hash of the row, so "home" is
  the winner only about half of the time and a model cannot learn the column order.
- `home_goals`/`away_goals`: sets won. Retirements keep the sets completed, with
  `status="finished"` and `finish_type="retired"` (every bet is void). Walkovers get
  `status="unavailable"`, `finish_type="walkover"` and no score.
- `odds` `{"1", "2"}`: the **average** closing prices (`AvgW`/`AvgL`, the Oddsportal market
  average published by tennis-data.co.uk, taken just before the start; the exact snapshot time
  is not documented). When they are missing, Pinnacle (`PSW`/`PSL`) and then Bet365
  (`B365W`/`B365L`) are used. The source hint `odds=avg|ps|b365` says which.
- `league`: `"ATP - SINGLES: {Tournament} ({Location}), {surface}"`, the FlashScore shape, so
  `sports.tennis.surface_of` and `best_of` work unchanged.
- `source`: `"tennis-data.co.uk;tour=atp;best_of=3;round=...;series=...;court=...;
  rank_home=..;rank_away=..;odds=avg;games=6-2 6-3"`; `extra(match)` parses it back. `games`
  lists every set's games from the home side.
- `kickoff`: the match date at 12:00 UTC. All games of one date share it, so with the 3-hour
  anti-leakage cutoff no game sees a result of its own day.
"""

import hashlib
import json
import re
import warnings
from datetime import date, datetime, timezone
from pathlib import Path

import httpx
from openpyxl import load_workbook

from footypreds.config import DATA
from footypreds.domain import Match

SOURCE = "tennis-data.co.uk"
BASE_URL = "https://www.tennis-data.co.uk/"
INDEX_URL = BASE_URL + "alldata.php"
# Directory seen on 2026-09-25; used only when the index page cannot be read.
KNOWN_PREFIX = "hrjk-85HytOjkhth76j_ygh4jf7/"
DATA_DIR = DATA / "benchmark" / "tennis"
RAW_DIR = DATA_DIR / "raw"
MANIFEST = DATA_DIR / "manifest.json"
TOURS = ("atp", "wta")
FIRST_YEAR = 2013  # first year published as .xlsx
LINK = re.compile(r"""href=["']?([^"'\s>]*?(\d{4})(w?)/\d{4}\.xlsx)""", re.IGNORECASE)
ODDS_COLUMNS = (("avg", "AvgW", "AvgL"), ("ps", "PSW", "PSL"), ("b365", "B365W", "B365L"))
SURFACES = {"hard": "hard", "clay": "clay", "grass": "grass", "carpet": "carpet"}


class TennisDataError(ValueError):
    pass


def discover(html):
    """{(tour, year): absolute URL} of every .xlsx workbook linked by the index page."""
    links = {}
    for href, year, women in LINK.findall(html or ""):
        url = href if href.startswith("http") else BASE_URL + href.lstrip("/")
        links[("wta" if women else "atp", int(year))] = url
    return links


def fallback_url(tour, year):
    return f"{BASE_URL}{KNOWN_PREFIX}{year}{'w' if tour == 'wta' else ''}/{year}.xlsx"


def raw_path(tour, year, directory=RAW_DIR):
    return Path(directory) / f"{tour}_{year}.xlsx"


def default_years(today=None):
    today = today or date.today()
    return list(range(FIRST_YEAR, today.year + 1))


def download(years=None, tours=TOURS, *, refresh=False, directory=RAW_DIR, client=None):
    """Fetch the workbooks that are missing (or all with `refresh`); returns manifest entries.

    Network use: only tennis-data.co.uk (public files, no key). Tests pass a `client` built on
    httpx.MockTransport. The current year's workbook grows during the season: pass
    `refresh=True` to update it.
    """
    years = list(years or default_years())
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    own = client is None
    client = client or httpx.Client(timeout=120, follow_redirects=True)
    manifest_path = directory.parent / MANIFEST.name
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        manifest = {}
    entries = []
    try:
        try:
            response = client.get(INDEX_URL)
            response.raise_for_status()
            links = discover(response.text)
        except httpx.HTTPError:
            links = {}
        for tour in tours:
            for year in years:
                path = raw_path(tour, year, directory)
                if path.exists() and not refresh:
                    continue
                url = links.get((tour, year)) or fallback_url(tour, year)
                response = client.get(url)
                if response.status_code == 404:
                    continue
                response.raise_for_status()
                content = response.content
                if not content.startswith(b"PK"):
                    raise TennisDataError(f"{url} nu este un fișier .xlsx.")
                path.write_bytes(content)
                entry = {
                    "tour": tour,
                    "year": year,
                    "url": url,
                    "file": path.name,
                    "bytes": len(content),
                    "sha256": hashlib.sha256(content).hexdigest(),
                    "downloaded_at": datetime.now(timezone.utc).isoformat(),
                }
                manifest[path.name] = entry
                entries.append(entry)
    finally:
        if own:
            client.close()
    manifest_path.write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")
    return entries


def number(value):
    """int/float cell -> float; blank, text or NaN -> None."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value) if value == value else None
    try:
        return float(str(value).strip().replace(",", "."))
    except ValueError:
        return None


def whole(value):
    value = number(value)
    return int(value) if value is not None and value == int(value) and value >= 0 else None


def day_of(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def home_is_winner(tour, day, winner, loser):
    """Deterministic, balanced orientation: a stable hash of the row decides."""
    key = f"{tour}|{day.isoformat()}|{winner}|{loser}".encode()
    return hashlib.sha256(key).digest()[0] % 2 == 0


def extra(match):
    """Source hints of a tennis-data match as a dict (strings), {} for other sources."""
    parts = match.source.split(";")
    if parts[0] != SOURCE:
        return {}
    return dict(part.split("=", 1) for part in parts[1:] if "=" in part)


def hint(value):
    return str(value).replace(";", ",").replace("=", "-").strip()


def odds_of(row):
    for name, winner, loser in ODDS_COLUMNS:
        w, lo = number(row.get(winner)), number(row.get(loser))
        if w and lo and 1 < w < 1001 and 1 < lo < 1001:
            return name, w, lo
    return None, None, None


def games_of(row):
    """[(winner_games, loser_games)] of the sets played."""
    sets = []
    for n in range(1, 6):
        w, lo = whole(row.get(f"W{n}")), whole(row.get(f"L{n}"))
        if w is None or lo is None:
            break
        sets.append((w, lo))
    return sets


def match_from_row(row, tour, today=None):
    """One workbook row -> Match; raises TennisDataError for unusable rows."""
    day = day_of(row.get("Date"))
    winner = str(row.get("Winner") or "").strip()
    loser = str(row.get("Loser") or "").strip()
    tournament = str(row.get("Tournament") or "").strip()
    if not (day and winner and loser and tournament) or winner == loser:
        raise TennisDataError("Rând incomplet.")
    if day > (today or date.today()):
        raise TennisDataError("Rezultat cu dată în viitor.")
    comment = str(row.get("Comment") or "Completed").strip().casefold()
    best = whole(row.get("Best of")) or 3
    if best not in (3, 5):
        raise TennisDataError("Număr de seturi necunoscut.")
    need = best // 2 + 1
    games = games_of(row)
    w_sets, l_sets = whole(row.get("Wsets")), whole(row.get("Lsets"))
    if w_sets is None or l_sets is None:
        w_sets = sum(w > lo for w, lo in games)
        l_sets = sum(lo > w for w, lo in games)
    finish, status = "", "finished"
    if comment.startswith("walk"):
        finish, status, w_sets, l_sets = "walkover", "unavailable", None, None
    elif comment.startswith(("retir", "disq", "def", "award")):
        # Retired, disqualified or awarded: the winner advanced, every bet is void.
        finish = "retired"
        if w_sets > need or l_sets > need:
            raise TennisDataError("Scor imposibil.")
    elif comment.startswith(("complet", "sched")):
        if best == 3 and w_sets == 3 and l_sets < 3:
            # A few Grand Slam rows say "Best of 3" for a five-set match.
            best, need = 5, 3
        if w_sets != need or l_sets >= need:
            raise TennisDataError("Scor final incoherent.")
    else:
        raise TennisDataError(f"Comentariu necunoscut: {comment}")

    home_wins = home_is_winner(tour, day, winner, loser)
    home, away = (winner, loser) if home_wins else (loser, winner)
    h_sets, a_sets = (w_sets, l_sets) if home_wins else (l_sets, w_sets)
    odds_source, w_odds, l_odds = odds_of(row)
    odds = {}
    if odds_source:
        odds = {"1": w_odds, "2": l_odds} if home_wins else {"1": l_odds, "2": w_odds}
    surface = SURFACES.get(str(row.get("Surface") or "").strip().casefold(), "")
    location = str(row.get("Location") or "").strip()
    name = f"{tournament} ({location})" if location and location != tournament else tournament
    league = f"{tour.upper()} - SINGLES: {name}"[:150]
    if surface:
        league = f"{league}, {surface}"
    w_rank, l_rank = whole(row.get("WRank")), whole(row.get("LRank"))
    ranks = (w_rank, l_rank) if home_wins else (l_rank, w_rank)
    hints = {
        "tour": tour,
        "best_of": best,
        "round": row.get("Round"),
        "series": row.get("Series") or row.get("Tier"),
        "court": row.get("Court"),
        "rank_home": ranks[0],
        "rank_away": ranks[1],
        "odds": odds_source,
    }
    if games:
        hints["games"] = " ".join(f"{w}-{lo}" if home_wins else f"{lo}-{w}" for w, lo in games)
    source = ";".join(
        [SOURCE] + [f"{k}={hint(v)}" for k, v in hints.items() if v not in (None, "")]
    )
    digest = hashlib.sha1(
        f"{day}|{tournament}|{row.get('Round')}|{winner}|{loser}".encode()
    ).hexdigest()[:12]
    return Match(
        id=f"td-{tour}-{day.year}-{digest}",
        kickoff=datetime(day.year, day.month, day.day, 12, tzinfo=timezone.utc),
        league=league,
        home=home[:120],
        away=away[:120],
        status=status,
        home_goals=h_sets,
        away_goals=a_sets,
        odds=odds,
        source=source,
        sport="tennis",
        finish_type=finish,
    )


def parse_workbook(source, tour, today=None):
    """(matches, rejected) of one workbook (a path or a binary file object)."""
    if tour not in TOURS:
        raise TennisDataError(f"Circuit necunoscut: {tour}")
    with warnings.catch_warnings():
        # openpyxl warns about an unknown workbook extension in these files; it is harmless.
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            return read_rows(workbook.worksheets[0].iter_rows(values_only=True), tour, today)
        finally:
            workbook.close()


def read_rows(rows, tour, today=None):
    """(matches, rejected) from worksheet value rows, the first being the header."""
    header = [str(cell).strip() if cell is not None else "" for cell in next(rows, ())]
    required = {"Date", "Winner", "Loser", "Tournament"}
    if not required.issubset(header):
        raise TennisDataError(f"Schema necunoscută pentru {tour}: lipsesc {required}.")
    matches, rejected, seen = [], 0, set()
    for values in rows:
        if not values or all(v is None for v in values):
            continue
        row = dict(zip(header, values))
        try:
            match = match_from_row(row, tour, today)
        except (TennisDataError, ValueError, TypeError):
            rejected += 1
            continue
        if match.id in seen:
            match = match.model_copy(update={"id": f"{match.id}-{len(seen)}"})
        seen.add(match.id)
        matches.append(match)
    return matches, rejected


def load_tennis_matches(years=None, tours=TOURS, directory=RAW_DIR, today=None):
    """Every stored workbook of `years` x `tours` as tennis Matches, oldest first.

    Walkovers are included (`status="unavailable"`); `HistoryIndex` skips them. Missing
    files are skipped; run `download()` (or `python -m footypreds.evaluation.tennis_eval
    --download`) first.
    """
    directory = Path(directory)
    if years is None:
        found = sorted(directory.glob("*_*.xlsx"))
        years = sorted({int(p.stem.split("_", 1)[1]) for p in found if p.stem[-4:].isdigit()})
    matches = []
    for tour in tours:
        for year in years:
            path = raw_path(tour, year, directory)
            if path.exists():
                matches.extend(parse_workbook(path, tour, today)[0])
    return sorted(matches, key=lambda m: (m.kickoff, m.id))
