import asyncio
import io
import time
from datetime import datetime, timedelta, timezone

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from footypreds.api import create_app, parse_csv
from footypreds.config import Settings
from footypreds.engine import analyze
from footypreds.engine.backtest import compact
from footypreds.provider import FlashScore, ProviderError, normalize_matches, parse_standings
from footypreds.store import Store
from footypreds.tests.helpers import fixture, fixtures_payload, h2h_payload, strong_history


def payload():
    return fixtures_payload()


def router(calls=None, standings=None):
    """Mock FlashScore: fixtures, h2h and standings endpoints."""

    def handle(request):
        if calls is not None:
            calls.append(request.url.path)
        path = request.url.path
        if path.endswith("matches/h2h"):
            kickoff = datetime.now(timezone.utc) + timedelta(days=1)
            return httpx.Response(200, json=h2h_payload(kickoff))
        if path.endswith("matches/standings"):
            return httpx.Response(200, json=standings or [])
        day = request.url.params.get("date")
        kickoff = None
        if day and day < datetime.now(timezone.utc).date().isoformat():
            kickoff = datetime.fromisoformat(day + "T15:00:00+00:00")
            data = fixtures_payload(kickoff, odds={})
            row = data[0]["matches"][0]
            row.update(
                match_id=f"past-{day}",
                match_status={"is_started": True, "is_finished": True},
                scores={"home": 2, "away": 1},
            )
            return httpx.Response(200, json=data)
        return httpx.Response(200, json=fixtures_payload(odds={"1": 1.5, "X": 4.2, "2": 6.5}))

    return httpx.MockTransport(handle)


@pytest.fixture
def client(tmp_path):
    settings = Settings(api_key="test-secret", database=tmp_path / "test.db")
    with TestClient(create_app(settings, router())) as test_client:
        yield test_client


def tomorrow():
    return (datetime.now(timezone.utc) + timedelta(days=1)).date().isoformat()


def test_normalizer_preserves_nulls_and_handles_bad_rows():
    data = fixtures_payload()
    data[0]["matches"].append({"match_id": "broken"})
    rows, rejected = normalize_matches(data)
    assert len(rows) == 1 and rejected == 1
    assert rows[0].home_goals is None
    assert rows[0].status == "scheduled"
    assert rows[0].odds == {"1": 1.5}


def test_results_parser_zero_is_valid():
    data = fixtures_payload()
    data[0]["matches"][0]["scores"] = {"home": 0, "away": 0}
    rows, _ = normalize_matches(data, results=True)
    assert rows[0].status == "finished" and rows[0].home_goals == 0


def test_h2h_parser_reads_flat_rows_with_string_scores():
    kickoff = datetime.now(timezone.utc)
    rows, rejected = normalize_matches(h2h_payload(kickoff), results=True)
    assert rejected == 0
    assert all(r.status == "finished" for r in rows)
    mutual = next(r for r in rows if r.id == "mutual")
    assert (mutual.home_goals, mutual.away_goals, mutual.home_id) == (3, 1, "")


def test_standings_parser_skips_bad_rows():
    table = parse_standings(
        [
            {"name": "A", "team_id": "a", "matches_played": 3, "goals": "7:2", "points": 9},
            {"name": "B", "goals": "bad"},
        ]
    )
    assert table == [
        {
            "position": 1,
            "team_id": "a",
            "name": "A",
            "played": 3,
            "wins": 0,
            "draws": 0,
            "losses": 0,
            "scored": 7,
            "conceded": 2,
            "points": 9,
        }
    ]


def test_unknown_schema_is_explicit_error():
    with pytest.raises(ProviderError):
        normalize_matches({"nonsense": []})


def test_home_health_static_and_secret_not_exposed(client):
    assert client.get("/").status_code == 200
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/app.css").status_code == 200
    assert client.get("/app.js").status_code == 200
    assert client.get("/.env").status_code == 404
    health = client.get("/api/health")
    assert health.json()["api_configured"]
    assert "test-secret" not in health.text
    assert "frame-ancestors 'none'" in health.headers["content-security-policy"]


def test_live_server_cors_only_allows_known_local_origins(client):
    allowed = client.get("/api/health", headers={"origin": "http://127.0.0.1:5500"})
    assert allowed.headers["access-control-allow-origin"] == "http://127.0.0.1:5500"
    blocked = client.get("/api/health", headers={"origin": "https://other.example"})
    assert "access-control-allow-origin" not in blocked.headers


def test_fixture_loading_and_cache(client):
    first = client.get(f"/api/matches?day={tomorrow()}").json()
    second = client.get(f"/api/matches?day={tomorrow()}").json()
    assert not first["cached"] and second["cached"]
    assert len(second["matches"]) == 1


def test_missing_match_invalid_params_and_csrf(client):
    assert client.post("/api/analyze/missing", json={}).status_code == 404
    assert client.get("/api/analysis/missing").status_code == 404
    assert client.get("/api/matches?day=bad").status_code == 422
    assert client.get("/api/demo?threshold=1.5").status_code == 422
    assert (
        client.post("/api/backtest", headers={"origin": "https://evil.example"}).status_code == 403
    )


def test_daily_board_predicts_every_fixture_without_history(client):
    board = client.get(f"/api/predictions?day={tomorrow()}").json()
    assert board["total"] == 1
    item = board["items"][0]
    p = item["probabilities"]
    assert p["1"] + p["X"] + p["2"] == pytest.approx(1)
    assert item["grade"] in "ABCD" and item["tip"]["label"]
    assert item["competition"] == "Test"
    assert client.get("/api/results").json()["metrics"]["selected"] == 0


def test_full_analysis_fetches_h2h_once_and_uses_every_competition(tmp_path):
    calls = []
    settings = Settings(api_key="k", database=tmp_path / "test.db")
    with TestClient(create_app(settings, router(calls))) as client:
        client.get(f"/api/matches?day={tomorrow()}")
        response = client.post("/api/analyze/fixture", json={"threshold": 0.6})
        body = response.json()
        assert response.status_code == 200
        assert calls.count("/api/flashscore/v2/matches/h2h") == 1
        prediction = body["prediction"]
        assert prediction["sample"]["home"] >= 6 and prediction["sample"]["away"] >= 6
        assert prediction["h2h"]["played"] == 1
        assert {"Cup", "Test"} <= {g["competition"] for g in prediction["form"]["home"]["last"]}
        assert all(item["id"] != "future" for item in prediction["form"]["home"]["last"])
        assert not body["retrospective"]
        # The cached h2h response is reused.
        client.post("/api/analyze/fixture", json={})
        assert calls.count("/api/flashscore/v2/matches/h2h") == 1


def test_started_match_analysis_is_retrospective_and_not_saved(client):
    store = client.app.state.store
    store.save_matches([fixture(kickoff=datetime.now(timezone.utc) - timedelta(hours=1))])
    body = client.post("/api/analyze/fixture", json={"enrich": False}).json()
    assert body["retrospective"] and not body["saved"]
    assert store.assessment_count() == 0


def test_excel_export_has_market_sheets(client):
    response = client.get(f"/api/export.xlsx?day={tomorrow()}")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == [
        "Legendă",
        "Predicții",
        "Scor corect",
        "Formă",
        "Valoare",
        "Pauză-Final",
    ]
    sheet = workbook["Predicții"]
    headers = [c.value for c in sheet[1]]
    row = dict(zip(headers, [c.value for c in sheet[2]]))
    assert row["Gazde"] == "Strong" and row["Cotă 1"] == 1.5
    assert row["1"] + row["X"] + row["2"] == pytest.approx(1)
    assert workbook["Formă"].max_row == 3


def test_history_sync_loads_past_days_once(tmp_path):
    calls = []
    settings = Settings(api_key="k", database=tmp_path / "test.db")
    with TestClient(create_app(settings, router(calls))) as client:
        started = client.post("/api/history/sync", json={"days": 3}).json()
        assert started["total"] == 3
        for _ in range(100):
            state = client.get("/api/history/sync").json()
            if state["status"] != "running":
                break
            time.sleep(0.01)
        assert state["status"] == "done" and state["matches"] == 3
        assert client.get("/api/health").json()["history_matches"] == 3
        fetched = len(calls)
        again = client.post("/api/history/sync", json={"days": 3}).json()
        # Yesterday is re-checked; older days are never fetched twice.
        assert again["total"] == 1
        assert client.post("/api/history/sync", json={"days": 91}).status_code == 422
        assert fetched == 3


def test_demo_never_persists_real_matches_or_predictions(client):
    assert client.get("/api/demo").json()["source"] == "synthetic"
    assert client.get(f"/api/predictions?day={tomorrow()}&demo=true").json()["total"] == 6
    assert client.app.state.store.matches() == []
    assert client.get("/api/results").json()["metrics"]["selected"] == 0


def test_snapshot_immutable_and_settlement_idempotent(tmp_path):
    store = Store(tmp_path / "ledger.db")
    match = fixture()
    p = compact(analyze(match, strong_history()))
    now = match.kickoff - timedelta(days=1)
    assert store.snapshot(match, p, now)
    assert not store.snapshot(match, p | {"threshold": 0.7}, now)
    final = match.model_copy(update={"status": "finished", "home_goals": 4, "away_goals": 0})
    assert store.settle([final]) == 1
    assert store.settle([final]) == 0
    assert store.predictions()[0]["prediction"]["threshold"] == 0.85
    assert store.predictions()[0]["result"]["won"]


def test_late_and_synthetic_predictions_not_saved(tmp_path):
    store = Store(tmp_path / "ledger.db")
    match = fixture()
    p = analyze(match, strong_history())
    assert not store.snapshot(match, p, match.kickoff)
    assert not store.snapshot(
        match.model_copy(update={"source": "synthetic"}), p, match.kickoff - timedelta(days=1)
    )
    assert store.assessment_count() == 0


def test_assessed_abstentions_count_towards_coverage(tmp_path):
    store = Store(tmp_path / "ledger.db")
    match = fixture()
    assert not store.snapshot(match, analyze(match, []), match.kickoff - timedelta(days=1))
    assert store.assessment_count() == 1
    assert store.predictions() == []


def test_store_cache_is_invalidated_by_writes(tmp_path):
    store = Store(tmp_path / "cache.db")
    assert store.matches() == []
    store.save_matches(strong_history()[:2])
    assert len(store.matches()) == 2
    first = store.matches()
    first.clear()
    assert len(store.matches()) == 2


def test_provider_429_has_no_retry_or_secret_leak(tmp_path):
    calls = []

    def handler(request):
        calls.append(request)
        return httpx.Response(429, json={"message": "do not expose token"})

    async def run():
        settings = Settings(api_key="private", database=tmp_path / "cache.db")
        provider = FlashScore(settings, Store(settings.database), httpx.MockTransport(handler))
        try:
            with pytest.raises(ProviderError) as exc:
                await provider.get("matches/list-by-date", {})
            assert exc.value.status == 429
            assert "private" not in str(exc.value)
            assert len(calls) == 1
        finally:
            await provider.client.aclose()

    asyncio.run(run())


CSV_HEADER = "id,kickoff,league,home,away,home_goals,away_goals\n"


def test_csv_rejects_missing_timezone_and_duplicates():
    with pytest.raises(ValueError):
        parse_csv(CSV_HEADER + "a,2025-01-01,League,A,B,1,0\n")
    row = "a,2025-01-01T18:00:00Z,League,A,B,1,0\n"
    with pytest.raises(ValueError):
        parse_csv(CSV_HEADER + row + row)
    assert len(parse_csv(CSV_HEADER + row)) == 1


def test_csv_endpoint_isolated_and_size_limited(client):
    response = client.post(
        "/api/backtest/csv", content=CSV_HEADER + "a,2025-01-01T18:00:00Z,League,A,B,1,0\n"
    )
    assert response.status_code == 200
    assert response.json()["metrics"]["accuracy"] is None
    assert client.app.state.store.matches() == []
    assert client.post("/api/backtest/csv", content="x" * 2_000_001).status_code == 413


def test_pre_1970_h2h_rows_do_not_crash_on_windows():
    rows, rejected = normalize_matches(
        [
            {
                "match_id": "old",
                "timestamp": -631152000,
                "tournament_name": "Friendly International",
                "home_team": {"name": "England"},
                "away_team": {"name": "Spain"},
                "scores": {"home": "2", "away": "1"},
            }
        ],
        results=True,
    )
    assert rejected == 0
    assert rows[0].kickoff.year == 1950
