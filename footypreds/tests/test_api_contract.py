"""HTTP contract of the FastAPI app: success/error paths, validation, security and isolation."""

import asyncio
import io
import itertools
import sqlite3
import time
import types
from datetime import datetime, timedelta, timezone

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import footypreds.provider as provider_module
from footypreds.api import create_app
from footypreds.competitions import match_competition, priority
from footypreds.config import Settings
from footypreds.domain import Match
from footypreds.engine import outcome
from footypreds.tests.helpers import h2h_payload, result

SECRET = "rapid-SECRET-9d8c7b6a"
SECURITY_HEADERS = {
    "x-content-type-options": "nosniff",
    "referrer-policy": "no-referrer",
}
LEAGUES = [
    ("ENGLAND: Premier League", "England"),
    ("SPAIN: LaLiga", "Spain"),
    ("ROMANIA: Liga 3", "Romania"),
    ("BHUTAN: Premier League", "Bhutan"),
    ("ITALY: Primavera U19", "Italy"),
    ("GERMANY: Regionalliga West", "Germany"),
]


POSTPONED = {n for n in range(30) if n % 13 == 5}
AVAILABLE = 30 - len(POSTPONED)


def today():
    return datetime.now(timezone.utc).date()


def day_rows(day, count=30):
    """A realistic day: several countries/tiers, some without odds, some postponed."""
    start = datetime.combine(day, datetime.min.time(), timezone.utc) + timedelta(hours=12)
    past = day < today()
    groups = {}
    for n in range(count):
        league, country = LEAGUES[n % len(LEAGUES)]
        row = {
            "match_id": f"{day.isoformat()}-{n:02}",
            "timestamp": (start + timedelta(minutes=15 * (n % 9))).timestamp(),
            "home_team": {"name": f"Home{n}" + (" W" if n % 11 == 0 else ""), "team_id": f"h{n}"},
            "away_team": {"name": f"Away{n}", "team_id": f"a{n}"},
            "match_status": {"is_started": past, "is_finished": past},
            "scores": {"home": n % 4, "away": n % 3} if past else {"home": None, "away": None},
            "odds": {"1": 1.8 + n % 3, "X": 3.4, "2": 4.1} if n % 4 else {},
        }
        if n % 13 == 5:
            row["status"] = "Postponed"
        groups.setdefault((league, country), []).append(row)
    return [{"name": k[0], "country_name": k[1], "matches": v} for k, v in groups.items()]


class FakeFlashScore:
    """Mock RapidAPI server. `fail` maps an endpoint suffix to a response factory."""

    def __init__(self):
        self.calls = []
        self.fail = {}
        self.delay = 0.0
        self.days = {}

    def __call__(self, request):
        return self.handle(request)

    async def handle(self, request):
        path = request.url.path
        self.calls.append((path, dict(request.url.params)))
        assert request.headers.get("x-rapidapi-key") == SECRET
        if self.delay:
            await asyncio.sleep(self.delay)
        for suffix, make in self.fail.items():
            if path.endswith(suffix):
                return make(request)
        if path.endswith("matches/h2h"):
            match_id = request.url.params["match_id"]
            kickoff = self.kickoffs.get(match_id, datetime.now(timezone.utc) + timedelta(days=1))
            rows = h2h_payload(kickoff)
            for row in rows:
                row["home_team"]["name"] = row["home_team"]["name"].replace("Strong", "Home1")
                row["away_team"]["name"] = row["away_team"]["name"].replace("Weak", "Away1")
            return httpx.Response(200, json=rows)
        if path.endswith("matches/standings"):
            return httpx.Response(200, json=[{"name": "Home1", "goals": "5:1", "points": 9}])
        day = datetime.fromisoformat(request.url.params["date"]).date()
        return httpx.Response(200, json=self.days.get(day) or day_rows(day))

    @property
    def kickoffs(self):
        out = {}
        for day, groups in self.days.items():
            for group in groups:
                for row in group["matches"]:
                    out[row["match_id"]] = datetime.fromtimestamp(row["timestamp"], timezone.utc)
        return out

    def count(self, suffix):
        return sum(path.endswith(suffix) for path, _ in self.calls)


@pytest.fixture(autouse=True)
def no_rate_limit_sleep(monkeypatch):
    counter = itertools.count(start=10_000, step=1_000)
    monkeypatch.setattr(
        provider_module, "time", types.SimpleNamespace(monotonic=lambda: next(counter))
    )


@pytest.fixture
def fake():
    return FakeFlashScore()


@pytest.fixture
def client(tmp_path, fake):
    settings = Settings(api_key=SECRET, database=tmp_path / "api.db")
    with TestClient(create_app(settings, httpx.MockTransport(fake))) as test_client:
        yield test_client


def tomorrow():
    return (today() + timedelta(days=1)).isoformat()


def wait_sync(client):
    for _ in range(500):
        state = client.get("/api/history/sync").json()
        if state["status"] != "running":
            return state
        time.sleep(0.01)
    raise AssertionError("sync did not finish")


# --- board paging ------------------------------------------------------------------------


def test_board_pages_partition_the_full_priority_ordered_list(client):
    full = client.get("/api/predictions", params={"day": tomorrow(), "limit": 400}).json()
    ids = [item["match"]["id"] for item in full["items"]]
    assert full["total"] == len(ids) == AVAILABLE  # postponed fixtures are excluded
    assert len(set(ids)) == len(ids)
    matches = [Match.model_validate(item["match"]) for item in full["items"]]
    assert matches == sorted(matches, key=priority)
    paged = []
    for offset in range(0, full["total"] + 7, 7):
        page = client.get(
            "/api/predictions", params={"day": tomorrow(), "limit": 7, "offset": offset}
        ).json()
        assert page["total"] == full["total"] and page["offset"] == offset
        assert len(page["items"]) <= 7
        paged += [item["match"]["id"] for item in page["items"]]
    assert paged == ids
    for item in full["items"]:
        p = item["probabilities"]
        assert p["1"] + p["X"] + p["2"] == pytest.approx(1)
        assert item["grade"] in "ABCD" and 0 <= item["confidence"] <= 100
        assert item["match"]["status"] != "unavailable"


def test_board_competition_filter_and_catalog(client):
    full = client.get("/api/predictions", params={"day": tomorrow(), "limit": 400}).json()
    catalog = {c["id"]: c["count"] for c in full["competitions"]}
    assert sum(catalog.values()) == full["total"]
    for competition, count in catalog.items():
        page = client.get(
            "/api/predictions",
            params={"day": tomorrow(), "competition": competition, "limit": 400},
        ).json()
        assert page["total"] == count
        assert all(
            match_competition(Match.model_validate(i["match"])) == competition
            for i in page["items"]
        )
    unknown = client.get("/api/predictions", params={"day": tomorrow(), "competition": "x|y"})
    assert unknown.status_code == 200 and unknown.json()["total"] == 0


def test_board_uses_the_fixture_cache(client, fake):
    for _ in range(3):
        client.get("/api/predictions", params={"day": tomorrow(), "limit": 5})
    assert fake.count("list-by-date") == 1
    client.get("/api/predictions", params={"day": tomorrow(), "refresh": True})
    assert fake.count("list-by-date") == 2


# --- validation bounds --------------------------------------------------------------------


@pytest.mark.parametrize(
    "method,url,kwargs",
    [
        ("get", "/api/predictions", {"params": {"day": "2026-13-01"}}),
        ("get", "/api/predictions", {"params": {"day": ""}}),
        ("get", "/api/predictions", {"params": {}}),
        ("get", "/api/predictions", {"params": {"day": "2026-06-01", "limit": 0}}),
        ("get", "/api/predictions", {"params": {"day": "2026-06-01", "limit": 401}}),
        ("get", "/api/predictions", {"params": {"day": "2026-06-01", "offset": -1}}),
        ("get", "/api/predictions", {"params": {"day": "2026-06-01", "limit": "ten"}}),
        ("get", "/api/analysis/x", {"params": {"threshold": 0.49}}),
        ("get", "/api/analysis/x", {"params": {"threshold": 1.0}}),
        ("get", "/api/analysis/x", {"params": {"threshold": "nan"}}),
        ("get", "/api/demo", {"params": {"threshold": 0.995}}),
        ("post", "/api/analyze/x", {"json": {"threshold": 0.3}}),
        ("post", "/api/analyze/x", {"json": {"threshold": "high"}}),
        (
            "post",
            "/api/analyze/x",
            {"content": b"{not json", "headers": {"content-type": "application/json"}},
        ),
        ("post", "/api/history/sync", {"json": {"days": 0}}),
        ("post", "/api/history/sync", {"json": {"days": 91}}),
        ("post", "/api/history/sync", {"json": {"days": 2.5}}),
        ("post", "/api/backtest", {"params": {"threshold": 0.2}}),
        ("post", "/api/plans", {"json": {"start_date": "2000-01-01"}}),
        ("post", "/api/plans", {"json": {"start_date": tomorrow(), "max_legs": 6}}),
        ("post", "/api/plans", {"json": {"start_date": tomorrow(), "competitions": ["no-bar"]}}),
        ("get", "/api/export.xlsx", {"params": {"day": "tomorrow"}}),
    ],
)
def test_invalid_parameters_are_422_json_without_network(client, fake, method, url, kwargs):
    response = getattr(client, method)(url, **kwargs)
    assert response.status_code == 422, response.text
    assert response.json()["detail"]
    assert fake.calls == []
    for header, value in SECURITY_HEADERS.items():
        assert response.headers[header] == value


def test_offset_past_the_end_is_an_empty_page(client):
    body = client.get("/api/predictions", params={"day": tomorrow(), "offset": 10_000}).json()
    assert body["items"] == [] and body["total"] == AVAILABLE


# --- error mapping ------------------------------------------------------------------------


@pytest.mark.parametrize(
    "response,status",
    [
        (lambda r: httpx.Response(429, text=SECRET), 429),
        (lambda r: httpx.Response(401, text=SECRET), 503),
        (lambda r: httpx.Response(403, text=SECRET), 503),
        (lambda r: httpx.Response(500, text=SECRET), 502),
        (lambda r: httpx.Response(200, text="<html>" + SECRET), 502),
        (lambda r: httpx.Response(200, json={"message": SECRET}), 502),
        (lambda r: httpx.Response(200, json={"unknown": "shape"}), 502),
    ],
)
def test_provider_failures_map_to_json_errors_without_the_secret(client, fake, response, status):
    fake.fail["list-by-date"] = response
    for url in ("/api/matches", "/api/predictions", "/api/export.xlsx"):
        reply = client.get(url, params={"day": tomorrow()})
        assert reply.status_code == status
        assert reply.headers["content-type"].startswith("application/json")
        assert reply.json()["detail"]
        assert SECRET not in reply.text and SECRET not in str(reply.headers)


def test_missing_api_key_is_503_and_never_calls_the_network(tmp_path):
    def refuse(request):
        raise AssertionError("network used without a key")

    settings = Settings(api_key="", database=tmp_path / "nokey.db")
    with TestClient(create_app(settings, httpx.MockTransport(refuse))) as client:
        assert client.get("/api/health").json()["api_configured"] is False
        assert client.get("/api/predictions", params={"day": tomorrow()}).status_code == 503
        assert client.get("/api/demo").status_code == 200


def test_unknown_routes_and_methods(client):
    assert client.get("/api/nothing").status_code == 404
    # GET on a POST-only route falls through to the static mount: never a success.
    assert client.get("/api/backtest").status_code in (404, 405)
    assert client.get("/api/plans/does-not-exist").status_code == 404
    assert client.post("/api/plans/does-not-exist/refresh").status_code == 404
    assert client.post("/api/analyze/missing", json={}).status_code == 404
    assert client.get("/api/analysis/missing").status_code == 404
    assert client.post("/api/backtest").status_code == 400


@pytest.mark.parametrize(
    "path",
    ["/.env", "/../.env", "/static/../../.env", "/%2e%2e/%2e%2e/.env", "/static/%2e%2e/api.py"],
)
def test_static_files_cannot_escape_the_web_folder(client, path):
    assert client.get(path).status_code == 404


# --- security headers, CSRF, CORS, hosts --------------------------------------------------


def assert_security_headers(response):
    for header, value in SECURITY_HEADERS.items():
        assert response.headers.get(header) == value, (response.status_code, header)
    assert "frame-ancestors 'none'" in response.headers.get("content-security-policy", "")


def test_security_headers_on_success_and_handled_errors(client):
    responses = [
        client.get("/"),
        client.get("/app.js"),
        client.get("/api/health"),
        client.get("/api/nothing"),
        client.get("/api/matches", params={"day": "bad"}),
        client.post("/api/analyze/missing", json={}),
        client.get("/api/health", headers={"host": "evil.example"}),
        client.get("/api/demo"),
    ]
    for response in responses:
        assert_security_headers(response)
    assert responses[-2].status_code == 400


def test_security_headers_on_rejected_cross_origin_post(client):
    response = client.post("/api/backtest", headers={"origin": "https://evil.example"})
    assert response.status_code == 403
    assert_security_headers(response)


def test_security_headers_on_unexpected_server_errors(tmp_path, fake):
    settings = Settings(api_key=SECRET, database=tmp_path / "boom.db")
    app = create_app(settings, httpx.MockTransport(fake))

    def explode():
        raise RuntimeError(f"internal detail {SECRET}")

    app.state.store.predictions = explode
    with TestClient(app, raise_server_exceptions=False) as client:
        response = client.get("/api/results")
        assert response.status_code == 500
        assert SECRET not in response.text and "internal detail" not in response.text
        assert response.json()["detail"]
        assert_security_headers(response)


@pytest.mark.parametrize(
    "origin,allowed",
    [
        ("https://evil.example", False),
        ("null", False),
        ("http://testserver.evil.example", False),
        ("https://testserver", False),
        ("http://evil.example:5500", False),
        ("http://testserver", True),
        ("http://127.0.0.1:5500", True),
        ("http://localhost:5501", True),
    ],
)
def test_csrf_origin_rules_for_post(client, origin, allowed):
    # An empty store answers 400: that proves the request reached the endpoint.
    response = client.post("/api/backtest", headers={"origin": origin})
    assert response.status_code == (400 if allowed else 403)


def test_post_without_origin_is_allowed_for_non_browser_clients(client):
    assert client.post("/api/backtest").status_code == 400


def test_cross_site_browser_get_cannot_spend_the_api_quota(client, fake):
    response = client.get(
        "/api/predictions", params={"day": tomorrow()}, headers={"sec-fetch-site": "cross-site"}
    )
    assert response.status_code == 403
    assert fake.calls == []
    same = client.get(
        "/api/predictions", params={"day": tomorrow()}, headers={"sec-fetch-site": "same-origin"}
    )
    assert same.status_code == 200


def test_cors_allowlist_and_preflight(client):
    preflight = {
        "origin": "http://localhost:5500",
        "access-control-request-method": "POST",
        "access-control-request-headers": "content-type",
    }
    ok = client.options("/api/analyze/x", headers=preflight)
    assert ok.status_code == 200
    assert ok.headers["access-control-allow-origin"] == "http://localhost:5500"
    evil = client.options("/api/analyze/x", headers=preflight | {"origin": "https://evil.example"})
    assert evil.status_code == 400
    assert "access-control-allow-origin" not in evil.headers
    wild = client.options(
        "/api/analyze/x", headers=preflight | {"access-control-request-method": "DELETE"}
    )
    assert wild.status_code == 400
    simple = client.get("/api/health", headers={"origin": "https://evil.example"})
    assert "access-control-allow-origin" not in simple.headers
    assert "access-control-allow-credentials" not in simple.headers


@pytest.mark.parametrize("host", ["evil.example", "127.0.0.2", "localhost.evil.example"])
def test_trusted_host_rejects_foreign_hosts(client, host):
    assert client.get("/api/health", headers={"host": host}).status_code == 400


def test_secret_never_appears_in_any_response(client, fake):
    fake.fail["matches/standings"] = lambda r: httpx.Response(500, text=SECRET)
    client.get("/api/predictions", params={"day": tomorrow()})
    match_id = f"{tomorrow()}-01"
    responses = [
        client.get("/api/health"),
        client.get("/openapi.json"),
        client.get("/docs"),
        client.get("/api/matches", params={"day": tomorrow()}),
        client.post(f"/api/analyze/{match_id}", json={"threshold": 0.6}),
        client.get(f"/api/analysis/{match_id}"),
        client.get("/api/results"),
        client.get("/api/plans"),
        client.get("/api/history/sync"),
        client.get("/api/export.xlsx", params={"day": tomorrow()}),
        client.get("/api/excel/health"),
    ]
    for response in responses:
        assert SECRET.encode() not in response.content, response.url
        assert SECRET not in str(response.headers)


# --- demo isolation -----------------------------------------------------------------------


def test_demo_endpoints_never_touch_the_store_or_network(tmp_path):
    def refuse(request):
        raise AssertionError("demo must not call FlashScore")

    settings = Settings(api_key=SECRET, database=tmp_path / "demo.db")
    with TestClient(create_app(settings, httpx.MockTransport(refuse))) as client:
        store = client.app.state.store
        version = store.version
        assert client.get("/api/demo", params={"threshold": 0.6}).json()["source"] == "synthetic"
        board = client.get("/api/predictions", params={"day": tomorrow(), "demo": True}).json()
        assert board["source"] == "synthetic" and board["total"] == 6
        assert client.get("/api/competitions", params={"day": tomorrow(), "demo": True}).json()
        assert client.post("/api/demo/backtest").json()["source"] == "synthetic"
        plan = client.post("/api/plans", json={"start_date": tomorrow(), "demo": True})
        assert plan.status_code == 202
        for _ in range(300):
            state = client.get(f"/api/plans/{plan.json()['id']}").json()
            if state["status"] != "generating":
                break
            time.sleep(0.01)
        assert state["source"] == "synthetic"
        assert store.matches() == [] and store.predictions() == []
        assert store.assessment_count() == 0 and store.version == version
        assert client.get("/api/results").json()["metrics"]["selected"] == 0


# --- analysis endpoints -------------------------------------------------------------------


def test_analysis_cache_is_invalidated_when_new_history_is_saved(client):
    client.get("/api/predictions", params={"day": tomorrow()})
    match_id = f"{tomorrow()}-01"
    first = client.get(f"/api/analysis/{match_id}").json()["prediction"]
    assert first["sample"]["home"] == 0
    store = client.app.state.store
    match = store.match(match_id)
    history = [
        result(f"hist{n}", 0, match.home, f"Opp{n}", 3, 0).model_copy(
            update={"kickoff": match.kickoff - timedelta(days=4 * n + 2)}
        )
        for n in range(10)
    ]
    store.save_matches(history)
    second = client.get(f"/api/analysis/{match_id}").json()["prediction"]
    assert second["sample"]["home"] == 10
    assert second != first
    # A changed fixture (new odds) is a different analysis even with the same history.
    store.save_matches([match.model_copy(update={"odds": {"1": 1.2, "X": 6.0, "2": 12.0}})])
    third = client.get(f"/api/analysis/{match_id}").json()["prediction"]
    assert third["components"]["market_1x2"] is not None
    assert third["components"]["market_1x2"] != second["components"]["market_1x2"]


def test_full_analysis_enriches_once_and_snapshots_at_the_ledger_threshold(client, fake):
    client.get("/api/predictions", params={"day": tomorrow()})
    match_id = f"{tomorrow()}-01"
    first = client.post(f"/api/analyze/{match_id}", json={"threshold": 0.5})
    assert first.status_code == 200
    body = first.json()
    assert not body["retrospective"] and body["standings"]
    assert body["prediction"]["sample"]["home"] > 0
    client.post(f"/api/analyze/{match_id}", json={"threshold": 0.9})
    assert fake.count("matches/h2h") == 1
    store = client.app.state.store
    assert store.assessment_count() == 1
    for row in store.predictions():
        assert row["prediction"]["threshold"] == 0.85
        assert row["prediction"]["selection"]["probability"] >= 0.85


def test_enrichment_quota_error_degrades_to_a_warning(client, fake):
    client.get("/api/predictions", params={"day": tomorrow()})
    fake.fail["matches/h2h"] = lambda r: httpx.Response(429)
    body = client.post(f"/api/analyze/{tomorrow()}-01", json={}).json()
    assert body["warnings"] and body["prediction"]["grade"] == "D"
    fake.fail["matches/h2h"] = lambda r: httpx.Response(401)
    assert client.post(f"/api/analyze/{tomorrow()}-01", json={}).status_code == 503


def test_started_or_finished_matches_are_retrospective(client):
    store = client.app.state.store
    past = datetime.now(timezone.utc) - timedelta(hours=2)
    store.save_matches(
        [
            Match(id="late", kickoff=past, league="L", home="A", away="B"),
            Match(
                id="done",
                kickoff=past,
                league="L",
                home="C",
                away="D",
                status="finished",
                home_goals=1,
                away_goals=0,
            ),
        ]
    )
    for match_id in ("late", "done"):
        body = client.post(f"/api/analyze/{match_id}", json={"enrich": False}).json()
        assert body["retrospective"] and not body["saved"]
        assert client.get(f"/api/analysis/{match_id}").json()["retrospective"]
    assert store.assessment_count() == 0


# --- export -------------------------------------------------------------------------------


def test_export_workbook_matches_the_board(client):
    board = client.get("/api/predictions", params={"day": tomorrow(), "limit": 400}).json()
    response = client.get("/api/export.xlsx", params={"day": tomorrow()})
    assert response.status_code == 200
    assert response.headers["content-disposition"] == (
        f'attachment; filename="FootyPreds-{tomorrow()}.xlsx"'
    )
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Predicții"]
    headers = [c.value for c in sheet[1]]
    rows = [dict(zip(headers, [c.value for c in row])) for row in sheet.iter_rows(min_row=2)]
    assert [(r["Gazde"], r["Oaspeți"]) for r in rows] == [
        (i["match"]["home"], i["match"]["away"]) for i in board["items"]
    ]
    for row, item in zip(rows, board["items"]):
        assert row["1"] == pytest.approx(item["probabilities"]["1"])
        assert row["Calitate"] == item["grade"]
    assert workbook["Formă"].max_row == 1 + 2 * len(rows)


def test_export_filters(client):
    only = client.get(
        "/api/export.xlsx", params={"day": tomorrow(), "competition": "england|premier league"}
    )
    rows = list(load_workbook(io.BytesIO(only.content))["Predicții"].iter_rows(min_row=2))
    assert len(rows) == sum(1 for n in range(0, 30, len(LEAGUES)) if n not in POSTPONED)
    assert (
        client.get(
            "/api/export.xlsx", params={"day": tomorrow(), "competition": "nowhere|none"}
        ).status_code
        == 404
    )
    # No history at all: every fixture is grade D, so the enriched-only export is empty.
    assert (
        client.get(
            "/api/export.xlsx", params={"day": tomorrow(), "enriched_only": True}
        ).status_code
        == 404
    )


# --- history sync -------------------------------------------------------------------------


def test_history_sync_is_idempotent_and_marks_only_final_days(client, fake):
    started = client.post("/api/history/sync", json={"days": 4})
    assert started.status_code == 202 and started.json()["total"] == 4
    state = wait_sync(client)
    assert state["status"] == "done" and state["done"] == 4
    store = client.app.state.store
    yesterday = today() - timedelta(days=1)
    assert yesterday.isoformat() not in store.synced_days()
    assert len(store.synced_days()) == 3
    finished = sum(m.status == "finished" for m in store.matches())
    assert finished == state["matches"] > 0
    calls = fake.count("list-by-date")
    again = client.post("/api/history/sync", json={"days": 4}).json()
    assert again["total"] == 1
    wait_sync(client)
    # Only yesterday is re-checked (possibly from the short-lived cache); older days never.
    assert fake.count("list-by-date") <= calls + 1
    assert sum(m.status == "finished" for m in store.matches()) == finished


def test_history_sync_quota_failure_keeps_completed_days(client, fake):
    seen = []

    def quota_on_third(request):
        seen.append(request)
        if len(seen) == 3:
            return httpx.Response(429)
        day = datetime.fromisoformat(request.url.params["date"]).date()
        return httpx.Response(200, json=day_rows(day, 4))

    fake.fail["list-by-date"] = quota_on_third
    client.post("/api/history/sync", json={"days": 5})
    state = wait_sync(client)
    assert state["status"] == "failed" and state["done"] == 2
    assert "Limita" in state["message"]
    store = client.app.state.store
    # Days 1 (yesterday, never marked) and 2 were fetched; only day 2 is final.
    assert store.synced_days() == {(today() - timedelta(days=2)).isoformat()}
    fake.fail.clear()
    retry = client.post("/api/history/sync", json={"days": 5}).json()
    assert retry["total"] == 4
    assert wait_sync(client)["status"] == "done"


def test_history_sync_unexpected_error_does_not_stay_running(client, monkeypatch):
    store = client.app.state.store

    def locked(matches):
        raise sqlite3.OperationalError("database is locked")

    monkeypatch.setattr(store, "save_matches", locked)
    client.post("/api/history/sync", json={"days": 2})
    state = wait_sync(client)
    assert state["status"] == "failed" and state["message"]


def test_history_sync_rejects_a_second_concurrent_run(client, fake):
    fake.delay = 0.05
    assert client.post("/api/history/sync", json={"days": 3}).status_code == 202
    busy = client.post("/api/history/sync", json={"days": 3})
    assert busy.status_code == 409 and busy.json()["detail"]
    assert wait_sync(client)["status"] == "done"


# --- CSV backtest -------------------------------------------------------------------------


CSV_HEADER = "id,kickoff,league,home,away,home_goals,away_goals\n"


def csv_rows(count, start=datetime(2025, 1, 1, 18, tzinfo=timezone.utc)):
    return "".join(
        f"c{n},{(start + timedelta(days=n // 6)).isoformat()},L,T{n % 6},U{n % 5},{n % 3},{n % 2}\n"
        for n in range(count)
    )


@pytest.mark.parametrize(
    "content,status",
    [
        (b"", 422),
        (b"id,kickoff\n1,2025-01-01T00:00:00Z\n", 422),
        (CSV_HEADER.encode(), 422),
        ((CSV_HEADER + "a,2025-01-01,L,A,B,1,0\n").encode(), 422),
        ((CSV_HEADER + "a,2025-01-01T18:00:00Z,L,A,A,1,0\n").encode(), 422),
        ((CSV_HEADER + "a,2025-01-01T18:00:00Z,L,A,B,-1,0\n").encode(), 422),
        ((CSV_HEADER + "a,2999-01-01T18:00:00Z,L,A,B,1,0\n").encode(), 422),
        (b"\xff\xfe\x00garbage", 422),
        ((CSV_HEADER + csv_rows(3001)).encode(), 422),
        (b"x" * 2_000_001, 413),
        (("﻿" + CSV_HEADER + csv_rows(12)).encode(), 200),
    ],
    ids=[
        "empty",
        "missing-columns",
        "header-only",
        "naive-kickoff",
        "same-teams",
        "negative-goals",
        "future",
        "not-utf8",
        "too-many-rows",
        "too-large",
        "bom-valid",
    ],
)
def test_csv_backtest_contract(client, content, status):
    response = client.post("/api/backtest/csv", content=content)
    assert response.status_code == status
    assert response.json().get("detail") or status == 200
    assert client.app.state.store.matches() == []


def test_past_day_board_shows_results_and_never_writes_the_ledger(client):
    day = (today() - timedelta(days=3)).isoformat()
    board = client.get("/api/predictions", params={"day": day, "limit": 400}).json()
    assert board["total"] == AVAILABLE
    for item in board["items"]:
        match = item["match"]
        assert match["status"] == "finished"
        assert item["result"]["score"] == f"{match['home_goals']}-{match['away_goals']}"
        assert item["result"]["tip_won"] == outcome(
            item["tip"]["key"], match["home_goals"], match["away_goals"]
        )
        # The fixture's own result is hidden from its prediction.
        assert item["sample"]["home"] == 0 and item["sample"]["away"] == 0
    match_id = board["items"][0]["match"]["id"]
    body = client.post(f"/api/analyze/{match_id}", json={"enrich": False}).json()
    assert body["retrospective"] and not body["saved"]
    assert client.app.state.store.assessment_count() == 0
