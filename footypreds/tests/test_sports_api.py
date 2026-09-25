"""Multi-sport HTTP API, store and competitions: sport params, enrichment, sync, export."""

import asyncio
import io
import itertools
import json
import time
import types
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import footypreds.provider as provider_module
from footypreds.api import board_item, create_app
from footypreds.competitions import catalog, competition_id, match_competition, priority
from footypreds.config import Settings
from footypreds.domain import Match
from footypreds.engine import summarize
from footypreds.engine.backtest import compact
from footypreds.provider import normalize_matches
from footypreds.sports import analyze_match, validate_analysis
from footypreds.store import Store
from footypreds.tests.helpers import fixture, fixtures_payload, h2h_payload

FIXTURES = Path(__file__).parent / "fixtures" / "flashscore"
SPORT_FILES = {"2": "tennis", "3": "basketball"}


def load(name):
    return json.loads((FIXTURES / name).read_text(encoding="utf-8"))


def today():
    return datetime.now(timezone.utc).date()


class Fake:
    """Mock FlashScore for every sport; records (path, params)."""

    def __init__(self):
        self.calls = []

    def __call__(self, request):
        path, params = request.url.path, dict(request.url.params)
        self.calls.append((path, params))
        if path.endswith("list-by-date"):
            sport = SPORT_FILES.get(params.get("sport_id"))
            if sport:
                return httpx.Response(200, json=load(f"list_{sport}.json"))
            return httpx.Response(200, json=fixtures_payload())
        if path.endswith("matches/live"):
            sport = SPORT_FILES.get(params.get("sport_id"), "football")
            return httpx.Response(200, json=load(f"live_{sport}.json"))
        match_id = params.get("match_id", "")
        sport = {"KnR6QDo1": "tennis", "KMHepeEM": "basketball"}.get(match_id)
        if path.endswith("matches/h2h"):
            if sport:
                return httpx.Response(200, json=load(f"h2h_{sport}.json"))
            return httpx.Response(200, json=h2h_payload(datetime.now(timezone.utc)))
        if path.endswith("matches/odds"):
            return httpx.Response(200, json=load(f"odds_{sport}.json") if sport else [])
        if path.endswith("matches/standings"):
            return httpx.Response(200, json=[])
        raise AssertionError(path)

    def count(self, suffix, **params):
        return sum(
            p.endswith(suffix) and all(q.get(k) == v for k, v in params.items())
            for p, q in self.calls
        )


@pytest.fixture(autouse=True)
def no_rate_limit_sleep(monkeypatch):
    counter = itertools.count(start=10_000, step=1_000)
    monkeypatch.setattr(
        provider_module, "time", types.SimpleNamespace(monotonic=lambda: next(counter))
    )


@pytest.fixture
def fake():
    return Fake()


@pytest.fixture
def client(tmp_path, fake):
    settings = Settings(api_key="k", database=tmp_path / "api.db")
    with TestClient(create_app(settings, httpx.MockTransport(fake))) as test_client:
        yield test_client


DAY = "2026-09-26"


def test_sports_endpoint_and_health(client):
    body = client.get("/api/sports").json()
    assert body["sports"] == [
        {"key": "football", "id": 1, "label": "Fotbal"},
        {"key": "basketball", "id": 3, "label": "Baschet"},
        {"key": "tennis", "id": 2, "label": "Tenis"},
    ]
    health = client.get("/api/health").json()
    assert health["sports"] == ["football", "basketball", "tennis"]
    assert set(health["history_by_sport"]) == {"football", "basketball", "tennis"}


@pytest.mark.parametrize(
    "url",
    [
        f"/api/predictions?day={DAY}&sport=curling",
        f"/api/matches?day={DAY}&sport=FOOTBALL",
        f"/api/competitions?day={DAY}&sport=",
        "/api/analysis/x?sport=golf",
    ],
)
def test_unknown_sport_is_a_422(client, fake, url):
    assert client.get(url).status_code == 422
    assert fake.calls == []


def test_matches_and_predictions_per_sport(client, fake):
    football = client.get(f"/api/matches?day={DAY}").json()
    assert football["sport"] == "football" and football["matches"][0]["sport"] == "football"
    board = client.get(f"/api/predictions?day={DAY}&sport=basketball").json()
    assert fake.count("list-by-date", sport_id="3") == 1
    assert board["sport"] == "basketball" and board["total"] == 30
    item = board["items"][0]
    assert item["sport"] == "basketball" and item["match"]["sport"] == "basketball"
    assert item["competition_id"].startswith("basketball:")
    assert [m["key"] for m in item["main"]][:2] == ["1", "2"]
    assert set(item["probabilities"]) == {m["key"] for m in item["main"]}
    assert "expected_goals" not in item and item["expected"]["home"] > 50
    assert all(c["id"].startswith("basketball:") for c in board["competitions"])
    tennis = client.get(f"/api/predictions?day={DAY}&sport=tennis&limit=400").json()
    ids = [i["competition_id"] for i in tennis["items"]]
    assert "tennis:atp - singles|chengdu" in ids
    # Main tours before Challenger, singles before doubles.
    first_doubles = next(
        n for n, i in enumerate(tennis["items"]) if "doubles" in i["competition_id"]
    )
    first_challenger = next(
        n for n, i in enumerate(tennis["items"]) if "challenger" in i["competition_id"]
    )
    singles = [n for n, i in enumerate(tennis["items"]) if "atp - singles" in i["competition_id"]]
    assert max(singles) < first_challenger < first_doubles


def test_football_board_keeps_its_fields(client):
    item = client.get(f"/api/predictions?day={DAY}").json()["items"][0]
    for key in ("expected_goals", "score", "scores", "htft", "probabilities", "tip", "form"):
        assert key in item
    assert item["sport"] == "football" and "over25" in item["probabilities"]
    assert [m["key"] for m in item["main"]] == ["1", "X", "2", "over25"]


def test_demo_board_is_football_only(client, fake):
    body = client.get(f"/api/predictions?day={DAY}&demo=true&sport=tennis").json()
    assert body["items"] == [] and body["source"] == "synthetic"
    assert fake.calls == []


def test_competitions_per_sport(client):
    local = client.get(f"/api/competitions?day={DAY}&sport=basketball").json()
    assert local["sport"] == "basketball"
    nba = next(c for c in local["competitions"] if c["id"] == "basketball:usa|nba")
    assert nba["popular"] and nba["count"] == 0
    refreshed = client.get(f"/api/competitions?day={DAY}&sport=tennis&refresh=true").json()
    ids = [c["id"] for c in refreshed["competitions"]]
    assert "tennis:atp - singles|wimbledon" in ids and ids[0].startswith("tennis:atp - singles|")


def test_tennis_enrichment_saves_history_and_merges_market_prices(client, fake):
    client.get(f"/api/predictions?day={DAY}&sport=tennis")
    store = client.app.state.store
    assert store.match("KnR6QDo1").odds == {"1": 2.3, "2": 1.57}
    assert client.post("/api/analyze/KnR6QDo1?sport=basketball", json={}).status_code == 404
    response = client.post("/api/analyze/KnR6QDo1?sport=tennis", json={"threshold": 0.6})
    assert response.status_code == 200
    body = response.json()
    assert fake.count("matches/odds", match_id="KnR6QDo1") == 1
    prediction = validate_analysis(body["prediction"])
    assert prediction["sport"] == "tennis" and prediction["sample"]["home"] > 10
    stored = store.match("KnR6QDo1")
    # List-by-date 1/2 stay; every other quoted market is added.
    assert stored.odds["1"] == 2.3 and stored.odds["2"] == 1.57
    assert stored.odds["ah_1_+1.5"] == 1.52 and stored.odds["sets_2-0"] == 4.0
    assert body["match"]["odds"]["ah_1_+1.5"] == 1.52
    by_key = {m["key"]: m for m in prediction["markets"]}
    assert by_key["ah_1_+1.5"]["odds"] == 1.52 and by_key["ah_1_+1.5"]["ev"] is not None
    history = store.matches(sport="tennis")
    assert any(m.finish_type == "retired" for m in history)
    assert store.matches(sport="football") == []
    # Reloading the day keeps the merged prices (odds union).
    client.get(f"/api/predictions?day={DAY}&sport=tennis&refresh=true")
    assert store.match("KnR6QDo1").odds["ah_1_+1.5"] == 1.52
    local = client.get("/api/analysis/KnR6QDo1?sport=tennis").json()
    assert local["prediction"]["sport"] == "tennis"


def test_basketball_enrichment_and_local_analysis(client, fake):
    client.get(f"/api/predictions?day={DAY}&sport=basketball")
    body = client.post("/api/analyze/KMHepeEM", json={}).json()
    prediction = validate_analysis(body["prediction"])
    assert prediction["sport"] == "basketball"
    assert fake.count("matches/standings", match_id="KMHepeEM") == 1
    assert client.get("/api/analysis/KMHepeEM?sport=tennis").status_code == 404
    assert client.get("/api/analysis/KMHepeEM").json()["prediction"]["sport"] == "basketball"


def test_tennis_enrichment_skips_standings(client, fake):
    client.get(f"/api/predictions?day={DAY}&sport=tennis")
    client.post("/api/analyze/KnR6QDo1", json={})
    assert fake.count("matches/standings") == 0


def wait_sync(client):
    for _ in range(200):
        state = client.get("/api/history/sync").json()
        if state["status"] != "running":
            return state
        time.sleep(0.02)
    raise AssertionError("sync did not finish")


def test_history_sync_per_sport(client, fake):
    started = client.post("/api/history/sync", json={"days": 3, "sports": ["basketball"]})
    assert started.status_code == 202 and started.json()["total"] == 3
    assert wait_sync(client)["status"] == "done"
    assert fake.count("list-by-date", sport_id="3") == 3
    assert fake.count("list-by-date", sport_id="1") == 0
    store = client.app.state.store
    older = {(today() - timedelta(days=n)).isoformat() for n in (2, 3)}
    assert store.synced_days("basketball") == older
    assert store.synced_days() == set() and store.synced_days("tennis") == set()
    # Default body: football only, as before.
    client.post("/api/history/sync", json={"days": 2})
    wait_sync(client)
    assert store.synced_days() == {(today() - timedelta(days=2)).isoformat()}
    assert fake.count("list-by-date", sport_id="1") == 2


@pytest.mark.parametrize("sports", [["curling"], [], ["tennis"] * 4])
def test_history_sync_rejects_bad_sports(client, sports):
    assert client.post("/api/history/sync", json={"sports": sports}).status_code == 422


def test_export_per_sport(client):
    response = client.get(f"/api/export.xlsx?day={DAY}&sport=basketball")
    assert response.status_code == 200
    assert "basketball" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Legendă", "Predicții", "Piețe"]
    assert workbook["Predicții"].max_row == 31
    football = client.get(f"/api/export.xlsx?day={DAY}")
    assert "Scor corect" in load_workbook(io.BytesIO(football.content)).sheetnames


def test_shared_state_objects_for_feature_routers(client):
    state = client.app.state
    for name in ("store", "provider", "cache", "settings", "enrich", "day_fixtures"):
        assert getattr(state, name) is not None
    assert state.excel_cache is state.cache and state.excel_enrich is state.enrich

    async def run():
        found, cached, rejected, settled = await state.day_fixtures(date(2026, 9, 26), "tennis")
        return found

    found = asyncio.run(run())
    assert found and {m.sport for m in found} == {"tennis"}


# --- store -----------------------------------------------------------------------------------


def test_store_odds_union_and_sport_filters(tmp_path):
    store = Store(tmp_path / "s.db")
    store.save_matches([fixture(odds={"1": 1.5, "X": 4.0, "2": 6.0})])
    store.save_matches([fixture(odds={"1": 1.5, "X": 4.0, "2": 6.0, "btts": 1.8})])
    store.save_matches([fixture(odds={"1": 1.6, "X": 3.9, "2": 5.5})])
    assert store.match("fixture").odds == {"1": 1.6, "X": 3.9, "2": 5.5, "btts": 1.8}
    other = fixture(id="b", sport="basketball", home_participant_id="p1")
    store.save_matches([other])
    store.save_matches([other.model_copy(update={"home_participant_id": ""})])
    assert store.match("b").home_participant_id == "p1"
    assert [m.id for m in store.matches(sport="basketball")] == ["b"]
    assert [m.id for m in store.matches_on(other.kickoff.date(), sport="football")] == ["fixture"]
    store.mark_synced(date(2026, 1, 1), 3)
    store.mark_synced(date(2026, 1, 2), 3, "tennis")
    assert store.synced_days() == {"2026-01-01"}
    assert store.synced_days("tennis") == {"2026-01-02"}


def tennis_pick(store, finish_type, home_sets=2, away_sets=0):
    kickoff = datetime.now(timezone.utc) + timedelta(days=1)
    match = Match(
        id=f"t-{finish_type or 'ok'}",
        kickoff=kickoff,
        league="ATP - SINGLES: X, hard",
        home="A",
        away="B",
        sport="tennis",
        odds={"1": 1.2, "2": 4.5},
    )
    prediction = analyze_match(match, [], 0.5)
    prediction["grade"], prediction["quality"] = "B", "sufficient"
    prediction["selection"] = next(m for m in prediction["markets"] if m["key"] == "1")
    assert store.snapshot(match, compact(prediction), datetime.now(timezone.utc))
    finished = match.model_copy(
        update={
            "status": "finished",
            "home_goals": home_sets,
            "away_goals": away_sets,
            "finish_type": finish_type,
        }
    )
    assert store.settle([finished]) == 1


def test_ledger_settles_every_sport_and_voids_retirements(tmp_path):
    store = Store(tmp_path / "s.db")
    tennis_pick(store, "")
    tennis_pick(store, "retired", 1, 0)
    rows = {r["match"]["id"]: r for r in store.predictions()}
    assert rows["t-ok"]["result"] == {"won": True, "score": "2-0"}
    assert rows["t-retired"]["result"]["won"] is None and rows["t-retired"]["result"]["void"]
    assert rows["t-ok"]["prediction"]["sport"] == "tennis"
    metrics = summarize(list(rows.values()), 2)
    assert metrics["settled"] == 1 and metrics["wins"] == 1
    assert metrics["void"] == 1 and metrics["pending"] == 0


def test_board_item_result_uses_universal_settlement():
    match = Match(
        id="g",
        kickoff=datetime(2026, 1, 1, tzinfo=timezone.utc),
        league="USA: NBA",
        country="USA",
        home="A",
        away="B",
        sport="basketball",
        status="finished",
        home_goals=101,
        away_goals=99,
    )
    item = board_item(match, analyze_match(match.model_copy(update={"odds": {"1": 1.3}}), []))
    assert item["result"]["score"] == "101-99"
    assert item["result"]["tip_won"] in (True, False, None)
    assert item["competition_id"] == "basketball:usa|nba"


# --- competitions ----------------------------------------------------------------------------


def test_football_competition_ids_are_unchanged():
    assert competition_id("ENGLAND: Premier League", "England") == "england|premier league"
    assert match_competition(fixture(league="SPAIN: LaLiga", country="Spain")) == "spain|laliga"


def test_sport_competition_ids_and_priority():
    tennis, _ = normalize_matches(load("list_tennis.json"), sport="tennis")
    ids = {match_competition(m) for m in tennis}
    assert "tennis:atp - singles|chengdu" in ids and "tennis:wta - doubles|seoul" in ids
    ordered = sorted(tennis, key=priority)
    categories = [m.league.split(":", 1)[0] for m in ordered]
    assert categories[0] in ("ATP - SINGLES", "WTA - SINGLES")
    assert categories.index("CHALLENGER MEN - SINGLES") < categories.index("ATP - DOUBLES")
    basketball, _ = normalize_matches(load("list_basketball.json"), sport="basketball")
    assert match_competition(basketball[0]) == "basketball:asia|asian games women"
    listed = catalog(basketball, sport="basketball")
    assert listed[0]["id"] == "basketball:usa|nba" and listed[0]["popular"]
    assert any(c["id"] == "basketball:australia|nbl" and c["popular"] for c in listed)
