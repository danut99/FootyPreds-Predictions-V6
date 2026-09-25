"""The daily .xlsx workbook (footypreds/excel.py build_workbook)."""

import io
import random
from datetime import date, timedelta, timezone

import pytest
from openpyxl import load_workbook

from footypreds.engine import analyze
from footypreds.excel import GRADE_FILLS, build_workbook
from footypreds.tests.generators import random_league, random_odds
from footypreds.tests.helpers import KICKOFF, fixture, strong_history

SHEETS = ["Legendă", "Predicții", "Scor corect", "Formă", "Valoare", "Pauză-Final"]


def items_for(count=6, seed=1):
    rng = random.Random(seed)
    teams = [f"Club{n}" for n in range(8)]
    history = random_league(rng, teams, 150, days=300)
    items = []
    for n in range(count):
        home, away = rng.sample(teams, 2)
        match = fixture(
            id=f"x{n}",
            home=home,
            away=away,
            league=f"ENGLAND: League {n % 2}",
            kickoff=KICKOFF + timedelta(minutes=10 * n),
            odds=random_odds(rng) if n % 2 else {},
        )
        items.append((match, analyze(match, history)))
    return items


def open_book(items, day=date(2026, 6, 1)):
    return load_workbook(io.BytesIO(build_workbook(day, items)))


def table(sheet):
    headers = [c.value for c in sheet[1]]
    return headers, [
        dict(zip(headers, [c.value for c in row])) for row in sheet.iter_rows(min_row=2)
    ]


def test_workbook_structure_and_values_match_the_analyses():
    items = items_for()
    book = open_book(items)
    assert book.sheetnames == SHEETS
    assert book.active.title == "Predicții"
    headers, rows = table(book["Predicții"])
    assert len(rows) == len(items)
    assert book["Predicții"].freeze_panes == "A2"
    for (match, analysis), row in zip(items, rows):
        p = {m["key"]: m["probability"] for m in analysis["markets"]}
        assert (row["Gazde"], row["Oaspeți"]) == (match.home, match.away)
        assert row["Competiție"] == match.league.split(":", 1)[1].strip()
        assert row["Ora (UTC)"] == match.kickoff.astimezone(timezone.utc).strftime("%H:%M")
        for header, key in (
            ("1", "1"),
            ("X", "X"),
            ("2", "2"),
            ("GG", "btts"),
            ("Peste 2.5", "over25"),
        ):
            assert row[header] == pytest.approx(p[key])
            assert 0 <= row[header] <= 1
        assert row["1"] + row["X"] + row["2"] == pytest.approx(1)
        assert row["Calitate"] == analysis["grade"]
        assert row["Încredere"] == analysis["confidence"]
        assert row["Cotă 1"] == match.odds.get("1")
        assert row["Scor probabil"] == analysis["scores"][0]["score"]
    percent = headers.index("1") + 1
    assert book["Predicții"].cell(row=2, column=percent).number_format == "0%"
    grade = headers.index("Calitate") + 1
    for n, (_, analysis) in enumerate(items, 2):
        fill = book["Predicții"].cell(row=n, column=grade).fill.fgColor.rgb
        assert fill.endswith(GRADE_FILLS[analysis["grade"]].fgColor.rgb[-6:])


def test_secondary_sheets():
    items = items_for()
    book = open_book(items)
    _, scores = table(book["Scor corect"])
    assert len(scores) == len(items)
    for (_, analysis), row in zip(items, scores):
        probs = [row[f"Prob. {n}"] for n in range(1, 6)]
        assert probs == sorted(probs, reverse=True)
        assert row["Scor 1"] == analysis["scores"][0]["score"]
    _, form = table(book["Formă"])
    assert len(form) == 2 * len(items)
    assert [r["Rol"] for r in form[:2]] == ["Gazde", "Oaspeți"]
    _, value = table(book["Valoare"])
    evs = [r["EV"] for r in value]
    assert evs == sorted(evs, reverse=True) and all(ev > 0 for ev in evs)
    expected = sum(1 for _, a in items for m in a["markets"] if m["ev"] is not None and m["ev"] > 0)
    assert len(value) == expected
    _, halves = table(book["Pauză-Final"])
    for (_, analysis), row in zip(items, halves):
        assert row["Pauză 1"] + row["Pauză X"] + row["Pauză 2"] == pytest.approx(1)
        assert row["HT/FT 1"] == analysis["htft"][0]["key"]
    legend = book["Legendă"]
    assert "2026-06-01" in legend["A1"].value
    assert legend["B3"].value == items[0][1]["version"]


def test_empty_day_still_builds_a_valid_workbook():
    book = open_book([])
    assert book.sheetnames == SHEETS
    assert book["Predicții"].max_row == 1
    assert book["Legendă"]["B3"].value == "—"


def test_no_history_rows_use_placeholders():
    match = fixture()
    book = open_book([(match, analyze(match, []))])
    _, rows = table(book["Predicții"])
    assert rows[0]["Formă gazde"] == "—" and rows[0]["Calitate"] == "D"
    _, form = table(book["Formă"])
    assert form[0]["Meciuri"] == 0 and form[0]["Puncte/meci"] is None


def test_unicode_names_round_trip():
    match = fixture(home="Steaua București ⭐", away="北京国安", league="ROMÂNIA: Liga Ⅰ")
    _, rows = table(open_book([(match, analyze(match, strong_history()))])["Predicții"])
    assert (rows[0]["Gazde"], rows[0]["Oaspeți"], rows[0]["Competiție"]) == (
        "Steaua București ⭐",
        "北京国安",
        "Liga Ⅰ",
    )


@pytest.mark.parametrize("name", ["+Club", "-Club", "@Club"])
def test_leading_operator_characters_stay_text(name):
    match = fixture(home=name)
    cell = open_book([(match, analyze(match, []))])["Predicții"]["C2"]
    assert cell.value == name and cell.data_type == "s"


def test_names_that_look_like_formulas_are_stored_as_text():
    match = fixture(home='=HYPERLINK("http://evil.example","Click")', league="=1+1")
    sheet = open_book([(match, analyze(match, []))])["Predicții"]
    assert sheet["C2"].data_type == "s" and sheet["B2"].data_type == "s"


def test_control_characters_in_names_do_not_break_the_export():
    match = fixture(home="Bad\x01Name", away="Tab\x0bTeam")
    book = open_book([(match, analyze(match, []))])
    assert "Name" in book["Predicții"]["C2"].value
